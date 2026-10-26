"""Konversi Markdown sederhana (judul, paragraf, daftar, tabel, gambar) ke DOCX.

Pemakaian: python ml/md_docx.py berkas.md [berkas2.md ...]  -> berkas.docx di folder yang sama
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Cm


def md_to_docx(text: str, base: Path, out: Path) -> None:
    doc = Document()
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                block.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            rows = [r for r in block if not set("".join(r)) <= {"-"}]
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for r, cells in enumerate(rows):
                for c, val in enumerate(cells):
                    par = t.cell(r, c).paragraphs[0]
                    add_runs(par, val)
                    if r == 0:  # baris judul tabel
                        for run in par.runs:
                            run.bold = True
            continue
        m = re.match(r"!\[.*?\]\((.+?)\)", line)
        if m:
            doc.add_picture(str(base / m.group(1)), width=Cm(15))
        elif line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            doc.add_heading(line.lstrip("# "), level=min(level, 3))
        elif line.lstrip().startswith("- "):
            style = "List Bullet 2" if line.startswith("  ") else "List Bullet"
            add_runs(doc.add_paragraph(style=style), line.lstrip()[2:])
        elif re.match(r"\d+\. ", line):
            # Nomor ditulis sebagai teks: gaya "List Number" python-docx melanjutkan nomor antardaftar
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Cm(0.63)
            par.paragraph_format.first_line_indent = Cm(-0.63)
            add_runs(par, line)
        elif line.strip():
            add_runs(doc.add_paragraph(), line)
        i += 1
    doc.save(out)


def add_runs(par, text: str) -> None:
    for part in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("*"):
            par.add_run(part[1:-1]).italic = True
        elif part.startswith("`"):
            par.add_run(part[1:-1]).font.name = "Consolas"
        else:
            par.add_run(part)


if __name__ == "__main__":
    for arg in sys.argv[1:]:
        src = Path(arg)
        md_to_docx(src.read_text(encoding="utf-8"), src.parent, src.with_suffix(".docx"))
        print(f"{src.with_suffix('.docx')}")
